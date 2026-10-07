#!/usr/bin/env python3
"""Read the supplied workbook without modifying it; preserve source cell provenance."""
import argparse
import hashlib
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INSTITUTION_PROFILE = json.loads((ROOT / 'public' / 'institution.json').read_text(encoding='utf-8'))


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def group_header(value):
    raw = clean(value)
    match = re.match(r"^(.*-\d{2})-(1|2)$", raw)
    return (match.group(1), int(match.group(2))) if match else (raw, None)


def subgroup_from_marks(marks):
    """The workbook repeats a subgroup lesson in the base column plus its subgroup column."""
    marks = set(marks)
    if 1 in marks and 2 not in marks:
        return 1
    if 2 in marks and 1 not in marks:
        return 2
    return None


def parse_lesson(raw, headers):
    text = clean(raw)
    # A shared class can begin with several group identifiers.
    while True:
        match = next((g for g in headers if text == g or text.startswith(g + " ")), None)
        if not match:
            break
        text = text[len(match):].lstrip()
    teacher_match = re.search(r"([А-ЯЁ][а-яё-]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.)", text)
    teacher = clean(teacher_match.group(1)) if teacher_match else ""
    before = text[:teacher_match.start()].strip() if teacher_match else text
    room = text[teacher_match.end():].strip() if teacher_match else ""
    kind_match = re.search(r"\s(ЛР|Л|П|К|Зачет|Зачёт|Экзамен|Консультация|Экз\.?|Зач\.?)$", before, re.I)
    kind = kind_match.group(1) if kind_match else ""
    subject = before[:kind_match.start()].strip() if kind_match else before
    return {"subject": subject, "type": kind, "teacher": teacher, "room": room, "raw": clean(raw)}


def import_schedule(path):
    workbook = load_workbook(path, data_only=True)
    schedule_profile = INSTITUTION_PROFILE.get("schedule", {})
    sheet_prefix = schedule_profile.get("sheetPrefix", "Расписание ФЗО")
    course_reference_year = int(schedule_profile.get("courseReferenceYear", datetime.now().year))
    sheets = [s for s in workbook if s.title.startswith(sheet_prefix)]
    raw_headers = sorted({clean(s.cell(1, c).value) for s in sheets for c in range(4, s.max_column + 1) if s.cell(1, c).value}, key=len, reverse=True)
    base_ids = sorted({group_header(header)[0] for header in raw_headers})
    groups_by_id = {}
    lessons = []
    lookup = {}
    audience_marks = {}
    occupied = 0
    unparsed = []

    for sheet in sheets:
        dates = sorted({sheet.cell(r, 1).value.date().isoformat() for r in range(2, sheet.max_row + 1) if isinstance(sheet.cell(r, 1).value, datetime)})
        for column in range(4, sheet.max_column + 1):
            source_group = clean(sheet.cell(1, column).value)
            if not source_group:
                continue
            group_id, subgroup = group_header(source_group)
            year = int(re.search(r"-(\d{2})", group_id).group(1))
            course = course_reference_year - (2000 + year) + 1
            group = groups_by_id.setdefault(group_id, {"id": group_id, "course": course, "dates": [], "lessons": [], "subgroups": {}})
            group["dates"] = sorted(set(group["dates"]) | set(dates))
            current_date = None

            for row in range(2, sheet.max_row + 1):
                value = sheet.cell(row, 1).value
                if isinstance(value, datetime):
                    current_date = value.date().isoformat()
                cell = sheet.cell(row, column)
                if not cell.value:
                    continue
                occupied += 1
                slot_match = re.search(r"\d+", clean(sheet.cell(row, 3).value))
                if not current_date or not slot_match:
                    raise ValueError(f"Missing date or pair at {sheet.title}!{cell.coordinate}")
                slot = int(slot_match.group())
                lines = [clean(line) for line in str(cell.value).splitlines() if clean(line)]

                for raw in lines:
                    key = (current_date, slot, raw)
                    if key not in lookup:
                        details = parse_lesson(raw, raw_headers)
                        index = len(lessons)
                        lookup[key] = index
                        lessons.append({"id": index, "date": current_date, "slot": slot, **details, "groups": [], "sources": []})
                        audience_marks[index] = {}
                        if not details["teacher"]:
                            unparsed.append({"sheet": sheet.title, "cell": cell.coordinate, "raw": raw})
                    index = lookup[key]
                    lesson = lessons[index]
                    if group_id not in lesson["groups"]:
                        lesson["groups"].append(group_id)
                    lesson["sources"].append({"sheet": sheet.title, "cell": cell.coordinate, "groupColumn": source_group})
                    audience_marks[index].setdefault(group_id, set()).add(subgroup or 0)
                    if index not in group["lessons"]:
                        group["lessons"].append(index)

    groups = list(groups_by_id.values())
    for group in groups:
        for index in group["lessons"]:
            subgroup = subgroup_from_marks(audience_marks[index].get(group["id"], set()))
            if subgroup:
                group["subgroups"][str(index)] = subgroup
        if not group["subgroups"]:
            group.pop("subgroups")

    groups.sort(key=lambda g: (g["course"], g["id"]))
    for lesson in lessons:
        lesson["groups"].sort(key=lambda g: g)
    date_match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", path.name)
    updated = f"{date_match.group(3)}-{date_match.group(2)}-{date_match.group(1)}" if date_match else "2026-10-05"
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    data = {
        "schema": 1,
        "version": updated + "-" + digest[:12] + "-g44",
        "updated": updated,
        "title": "Осенняя установочная сессия 2026/27",
        "timezone": INSTITUTION_PROFILE.get("institution", {}).get("timezone", "Europe/Moscow"),
        "institution": {"id": INSTITUTION_PROFILE.get("institution", {}).get("id", "unknown"), "unit": INSTITUTION_PROFILE.get("unit", {}).get("id", "unknown")},
        "source": {"file": path.name, "sha256": digest, "adapter": schedule_profile.get("sourceAdapter", "xlsx")},
        "groups": groups,
        "lessons": lessons,
    }
    subgroup_lessons = sum(len(group.get("subgroups", {})) for group in groups)
    report = {
        "groups": len(groups),
        "sourceGroupColumns": len(raw_headers),
        "subgroupFamilies": len([base for base in base_ids if base + "-1" in raw_headers or base + "-2" in raw_headers]),
        "subgroupLessons": subgroup_lessons,
        "uniqueLessons": len(lessons),
        "sourceCells": occupied,
        "courses": sorted({g["course"] for g in groups}),
        "unparsed": unparsed,
        "adapter": schedule_profile.get("sourceAdapter", "xlsx"),
        "institution": INSTITUTION_PROFILE.get("institution", {}).get("id", "unknown"),
        "unit": INSTITUTION_PROFILE.get("unit", {}).get("id", "unknown"),
    }
    assert {g["id"] for g in groups} == set(base_ids), "Base groups lost or duplicated"
    assert not any(re.search(r"-[12]$", g["id"]) for g in groups), "Technical subgroup leaked into group selector"
    assert sum(len(l["sources"]) for l in lessons) >= occupied, "Source records lost"
    assert all(l["subject"] and 1 <= l["slot"] <= 7 for l in lessons)
    return data, report

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("public/schedule.json"))
    args = parser.parse_args()
    data, report = import_schedule(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    full_path = Path("data/schedule-full.json")
    full_path.parent.mkdir(parents=True, exist_ok=True)
    full_path.write_text(text, encoding="utf-8")
    # Provenance stays in the source package; ship only fields used by the app.
    public_data = {k: v for k, v in data.items() if k != "source"}
    public_data["lessons"] = [{k: v for k, v in lesson.items() if k not in ("groups", "sources", "raw")} for lesson in data["lessons"]]
    public_text = json.dumps(public_data, ensure_ascii=False, separators=(",", ":"))
    args.output.write_text(public_text, encoding="utf-8")
    args.output.with_suffix(".js").write_text("window.RGATU_DATA=" + public_text + ";\n", encoding="utf-8")
    Path("import-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({**report, "unparsed": report["unparsed"][:12], "jsonBytes": len(public_text.encode())}, ensure_ascii=False))


if __name__ == "__main__":
    main()
